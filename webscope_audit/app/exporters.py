from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _flatten_report(scan_result: dict) -> list[tuple[str, str]]:
    rows = [
        ("Cible", str(scan_result.get("target", ""))),
        ("Date", str(scan_result.get("collected_at", ""))),
        ("Périmètre", str(scan_result.get("scope", ""))),
        ("Titre", str(scan_result.get("title", "WebScope Audit V4"))),
    ]
    rows.extend([
        ("DNS", str(scan_result.get("dns", {}))),
        ("HTTP", str(scan_result.get("http", {}))),
        ("TLS", str(scan_result.get("tls", {}))),
        ("Technologies", str(scan_result.get("technologies", {}))),
        ("Sous-domaines", ", ".join(scan_result.get("subdomains", [])) or "Aucun"),
        ("Ports ouverts", ", ".join(str(item) for item in scan_result.get("ports", {}).get("open", [])) or "Aucun"),
        ("Erreurs", "; ".join(scan_result.get("errors", [])) or "Aucune"),
        ("Recommandations", "; ".join(scan_result.get("recommendations", [])) or "Aucune"),
    ])
    return rows


def export_docx_report(scan_result: dict, output_path: str | Path) -> Path:
    path = Path(output_path)
    document = Document()
    document.add_heading(scan_result.get("title", "WebScope Audit V4"), level=1)
    document.add_paragraph(f"Cible : {scan_result.get('target', '')}")
    document.add_paragraph(f"Date : {scan_result.get('collected_at', '')}")
    document.add_paragraph(f"Périmètre : {scan_result.get('scope', '')}")

    for title, value in _flatten_report(scan_result):
        p = document.add_paragraph()
        p.add_run(f"{title}: ").bold = True
        p.add_run(str(value)[:4000])

    document.save(path)
    return path


def export_pdf_report(scan_result: dict, output_path: str | Path) -> Path:
    path = Path(output_path)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleStyle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, textColor=colors.HexColor("#0f172a"), leading=20)
    heading_style = ParagraphStyle("HeadingStyle", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor("#1d4ed8"))
    body_style = ParagraphStyle("BodyStyle", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12)

    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    story = []
    story.append(Paragraph(scan_result.get("title", "WebScope Audit V4"), title_style))
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"Cible : {scan_result.get('target', '')}", body_style))
    story.append(Paragraph(f"Date : {scan_result.get('collected_at', '')}", body_style))
    story.append(Paragraph(f"Périmètre : {scan_result.get('scope', '')}", body_style))
    story.append(Spacer(1, 12))

    rows = [["Section", "Détail"]]
    for label, value in _flatten_report(scan_result):
        rows.append([label, (str(value)[:180] or "-")])
    table = Table(rows, colWidths=[120, 350])
    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbeafe")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ])
    )
    story.append(table)
    doc.build(story)
    return path
